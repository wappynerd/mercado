"""Baja el ITCRM diario del BCRA y lo deja en data/itcrm.json para que lo lea la página.

El BCRA lo publica como Excel, y el navegador no puede leerlo directo (no permite pedidos
desde otra página). Esta tarea corre una vez por día en GitHub Actions.

Uso: python scripts/itcrm.py
"""

from __future__ import annotations

import datetime
import io
import json
import pathlib
import urllib.request

import openpyxl

URL = "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/ITCRMSerie.xlsx"
DESDE = "2015-12-17"   # base del índice: 17-12-2015 = 100, y salida del cepo anterior
SALIDA = pathlib.Path(__file__).resolve().parent.parent / "data" / "itcrm.json"


def bajar() -> bytes:
    pedido = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(pedido, timeout=120) as respuesta:
        return respuesta.read()


def serie(contenido: bytes) -> list[list]:
    """Devuelve [[fecha ISO, ITCRM], ...] de la primera hoja (serie diaria)."""
    libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro[libro.sheetnames[0]]
    filas = []
    for fila in hoja.iter_rows(values_only=True):
        fecha, valor = fila[0], fila[1]
        if isinstance(fecha, datetime.datetime) and isinstance(valor, (int, float)):
            iso = fecha.date().isoformat()
            if iso >= DESDE:
                filas.append([iso, round(float(valor), 3)])
    return filas


def main() -> None:
    filas = serie(bajar())
    if len(filas) < 1000:
        raise SystemExit(f"Serie sospechosamente corta ({len(filas)} filas): no se pisa el archivo.")
    SALIDA.parent.mkdir(exist_ok=True)
    SALIDA.write_text(
        json.dumps({"fuente": "BCRA, ITCRM diario (17-12-2015 = 100)", "hasta": filas[-1][0], "serie": filas},
                   separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"{len(filas)} filas, hasta {filas[-1][0]}: ITCRM {filas[-1][1]}")


if __name__ == "__main__":
    main()
